"""Read spreadsheet files without changing their source bytes."""
from __future__ import annotations
import re
from pathlib import Path
from dataclasses import dataclass, field
from datetime import datetime, date
import numpy as np
import pandas as pd
import openpyxl

SUPPORTED = {'.xlsx', '.xlsm', '.xls', '.csv'}

def safe_identifier(value: str) -> str:
    value = re.sub(r'[^a-zA-Z0-9_]+', '_', str(value)).strip('_').lower() or 'column'
    return ('t_' if value[0].isdigit() else '') + value

def clean_value(value):
    if value is None: return None
    try:
        if pd.isna(value): return None
    except (TypeError, ValueError): pass
    if isinstance(value, (datetime, date, pd.Timestamp)): return value.isoformat()
    if isinstance(value, np.generic): return value.item()
    if isinstance(value, float) and not np.isfinite(value): return None
    return value

@dataclass
class Table:
    name: str
    alias: str
    file: str
    sheet: str
    frame: pd.DataFrame
    columns: list[dict]
    warnings: list[str] = field(default_factory=list)
    header_row: int = 1
    def metadata(self):
        return dict(name=self.name, alias=self.alias, file=self.file, sheet=self.sheet,
                    row_count=len(self.frame), columns=self.columns, warnings=self.warnings,
                    header_row=self.header_row)

def _header(rows):
    candidates = []
    for i, row in enumerate(rows[:30]):
        present = [v for v in row if v is not None and str(v).strip()]
        if not present: continue
        text = sum(isinstance(v,str) and not str(v).startswith('=') for v in present)
        # A title has few filled cells; headers are wide, unique, and textual.
        score = len(present)*2 + text + len(set(str(v) for v in present))/len(present)
        if text == len(present): score += 2
        candidates.append((score, -i, i))
    return max(candidates)[2] if candidates else 0

def _infer(series, name, formats):
    nonnull = series.dropna()
    if nonnull.empty: return series, 'text'
    if all(isinstance(v,(bool,np.bool_)) for v in nonnull): return series, 'boolean'
    if pd.api.types.is_datetime64_any_dtype(series) or all(isinstance(v,(datetime,date)) for v in nonnull):
        return pd.to_datetime(series,errors='coerce'), 'date'
    if pd.api.types.is_numeric_dtype(series):
        kind = 'percent' if (any('%' in f for f in formats) or bool(re.search(r'pct|percent',name,re.I))) else 'currency' if (any('$' in f or '€' in f or '£' in f for f in formats) or bool(re.search(r'usd|price|cost|salary',name,re.I))) else 'number'
        return pd.to_numeric(series,errors='coerce'), kind
    stripped = nonnull.astype(str).str.strip()
    if stripped.str.lower().isin(['true','false','yes','no']).all():
        return series.map(lambda v: None if pd.isna(v) else str(v).lower() in ('true','yes')), 'boolean'
    if re.search(r'date|time|birthday|dob',name,re.I) or stripped.str.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}').mean()>.9:
        parsed = pd.to_datetime(series,errors='coerce',format='mixed')
        if parsed.notna().sum() >= .9*len(nonnull): return parsed,'date'
    numeric = pd.to_numeric(series.astype(str).str.replace(r'[$€£,%\s,]','',regex=True),errors='coerce')
    # Keep IDs and postal/phone strings as text, particularly leading zeros.
    if numeric.notna().sum()==len(nonnull) and not re.search(r'id$|sku|phone|postal|zip',name,re.I):
        kind='percent' if stripped.str.contains('%',regex=False).any() else 'currency' if stripped.str.contains(r'[$€£]').any() else 'number'
        if kind=='percent': numeric/=100
        return numeric,kind
    return series.map(lambda v: None if pd.isna(v) else str(v)), 'text'

def _table(path, sheet, rows, formulas=None, formats=None):
    rows = [list(r) for r in rows]
    if not rows or not any(any(v is not None for v in r) for r in rows): return None
    h = _header(rows)
    width=max(len(r) for r in rows)
    rows=[r+[None]*(width-len(r)) for r in rows]
    active=[j for j in range(width) if rows[h][j] is not None or any(r[j] is not None for r in rows[h+1:])]
    names=[]; originals=[]; used=set()
    for j in active:
        original=str(rows[h][j]).strip() if rows[h][j] is not None else f'Column {j+1}'
        safe=safe_identifier(original); unique=safe; suffix=2
        while unique in used: unique=f'{safe}_{suffix}'; suffix+=1
        used.add(unique); names.append(unique); originals.append(original)
    body=[[r[j] for j in active] for r in rows[h+1:] if any(r[j] is not None for j in active)]
    frame=pd.DataFrame(body,columns=names)
    columns=[]; warnings=[]
    for name,original,j in zip(names,originals,active):
        uncalculated=0
        if formulas:
            uncalculated=sum(1 for i in range(h+1,len(rows)) if i<len(formulas) and j<len(formulas[i]) and isinstance(formulas[i][j],str) and formulas[i][j].startswith('=') and rows[i][j] is None)
        fmt=[formats.get((i+1,j+1),'') for i in range(h+1,min(len(rows),h+50))] if formats else []
        frame[name],kind=_infer(frame[name],original,fmt)
        s=frame[name]; unique=s.dropna().unique()
        column=dict(name=name,original_name=original,type=kind,null_pct=round(float(s.isna().mean()*100),2) if len(s) else 0,
                    distinct_count=len(unique),sample_values=[clean_value(v) for v in unique[:5]],uncalculated=bool(uncalculated),uncalculated_count=uncalculated,min=None,max=None,distinct_values=None)
        if kind in ('number','currency','percent','date') and len(unique):
            column.update(min=clean_value(s.min()),max=clean_value(s.max()))
        if kind in ('text','boolean') and len(unique)<=50: column['distinct_values']=[clean_value(v) for v in unique]
        columns.append(column)
        if uncalculated: warnings.append(f'{original}: Some formula results were not saved. Open and save the workbook, then reload it to include them.')
    name=safe_identifier(path.stem)+'__'+safe_identifier(sheet)
    return Table(name,f'{path.stem.title()} · {sheet}',path.name,sheet,frame,columns,warnings,h+1)

def ingest_file(path):
    path=Path(path)
    if path.suffix.lower() not in SUPPORTED: raise ValueError('Unsupported file. Use .xlsx, .xlsm, .xls, or .csv.')
    tables=[]
    if path.suffix.lower()=='.csv':
        try: raw=pd.read_csv(path,header=None,dtype=object,encoding='utf-8-sig')
        except UnicodeDecodeError: raw=pd.read_csv(path,header=None,dtype=object,encoding='latin-1')
        table=_table(path,'Sheet1',raw.where(raw.notna(),None).values.tolist())
        if table: tables.append(table)
    elif path.suffix.lower()=='.xls':
        for sheet,raw in pd.read_excel(path,sheet_name=None,header=None).items():
            table=_table(path,sheet,raw.where(raw.notna(),None).values.tolist())
            if table: tables.append(table)
    else:
        values=openpyxl.load_workbook(path,data_only=True,read_only=False,keep_links=False)
        formulas=openpyxl.load_workbook(path,data_only=False,read_only=False,keep_links=False)
        try:
            for ws in values:
                rows=[list(r) for r in ws.values]; fr=[list(r) for r in formulas[ws.title].values]
                # Repeat merged labels, preserving source workbook.
                for region in ws.merged_cells.ranges:
                    value=ws.cell(region.min_row,region.min_col).value
                    for i in range(region.min_row-1,region.max_row):
                        for j in range(region.min_col-1,region.max_col): rows[i][j]=value
                fm={(c.row,c.column):c.number_format for row in ws for c in row if c.value is not None}
                table=_table(path,ws.title,rows,fr,fm)
                if table: tables.append(table)
        finally: values.close(); formulas.close()
    return tables

def ingest_directory(folder):
    tables={}; errors=[]
    for path in sorted(Path(folder).glob('*')):
        if path.name.startswith('.') or path.suffix.lower() not in SUPPORTED: continue
        try:
            for table in ingest_file(path):
                if table.name in tables: raise ValueError('Table name collides with another file or sheet.')
                tables[table.name]=table
        except Exception: errors.append({'file':path.name,'message':'Could not read this workbook. Check that it is a valid, unlocked spreadsheet.'})
    return tables,errors
